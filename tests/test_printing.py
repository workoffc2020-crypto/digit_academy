import os
from datetime import date

import pytest

from app.printing import docs, export, render
from app.services import attendance as att
from app.services import fees, reports, students
from app.utils import money

fitz = pytest.importorskip("pymupdf")


@pytest.fixture()
def paid(ctx, admin, student, course_id):
    eid = fees.enroll_with_fees(ctx.db, admin, student, course_id, course_fee=1500000, admission_fee=100000, installments=3, first_due=date(2026, 5, 1))
    pid = fees.collect_payment(ctx.db, admin, student, 650000, "Cash", enrollment_id=eid)
    return eid, pid


def _text(path):
    d = fitz.open(path)
    t = "".join(p.get_text() for p in d)
    n = len(d)
    d.close()
    return t, n


def test_receipt_a4_pdf_has_saved_values_and_logo(ctx, admin, paid, tmp_path):
    eid, pid = paid
    p = render.to_pdf(docs.receipt(ctx.db, pid, "A4"), str(tmp_path / "r.pdf"))
    t, n = _text(p)
    no = fees.get_payment(ctx.db, pid)["receipt_no"]
    assert no in t and "Asha Verma" in t and "6,500.00" in t
    assert "Six Thousand Five Hundred" in t and "Test Academy" in t
    d = fitz.open(p)
    assert d[0].get_images(), "logo must be embedded"


def test_receipt_thermal_is_narrow_single_page(ctx, admin, paid, tmp_path):
    p = render.to_pdf(docs.receipt(ctx.db, paid[1], "THERMAL80"), str(tmp_path / "t.pdf"))
    d = fitz.open(p)
    assert len(d) == 1 and abs(d[0].rect.width - 80 / 25.4 * 72) < 2 and d[0].rect.height > 150


def test_reversed_receipt_is_marked_and_reversal_note_references_it(ctx, admin, paid, tmp_path):
    rid = fees.reverse_payment(ctx.db, admin, paid[1], "wrong entry")
    t1, _ = _text(render.to_pdf(docs.receipt(ctx.db, paid[1]), str(tmp_path / "a.pdf")))
    t2, _ = _text(render.to_pdf(docs.receipt(ctx.db, rid), str(tmp_path / "b.pdf")))
    assert "CANCELLED" in t1 and "REVERSAL" in t2 and fees.get_payment(ctx.db, paid[1])["receipt_no"] in t2


def test_admission_form_profile_and_id_card(ctx, admin, student, paid, tmp_path):
    t, _ = _text(render.to_pdf(docs.admission_form(ctx.db, ctx.vault, admin, student), str(tmp_path / "f.pdf")))
    assert "ADMISSION FORM" in t and "Ram Verma" in t and "Python" in t
    t, _ = _text(render.to_pdf(docs.admission_form(ctx.db, ctx.vault, admin, student, True), str(tmp_path / "p.pdf")))
    assert "Fee summary" in t and "Outstanding" in t
    p = render.to_pdf(docs.id_cards(ctx.db, ctx.vault, admin, [student, student]), str(tmp_path / "c.pdf"))
    d = fitz.open(p)
    assert len(d) == 2 and abs(d[0].rect.width - 85.6 / 25.4 * 72) < 2


def test_hindi_names_are_embedded_with_a_devanagari_font(ctx, admin, tmp_path):
    sid = students.create(ctx.db, admin, ctx.vault, {"name": "आशा वर्मा", "mobile": "9876501234"})
    p = render.to_pdf(docs.admission_form(ctx.db, ctx.vault, admin, sid), str(tmp_path / "h.pdf"))
    d = fitz.open(p)
    assert any("FreeSans" in f[3] for f in d[0].get_fonts())
    assert "आशा" in d[0].get_text() or "आ" in d[0].get_text()


def test_report_tables_and_exports(ctx, admin, paid, tmp_path):
    ds = reports.run(ctx.db, admin, "outstanding")
    t, _ = _text(render.to_pdf(docs.dataset(ctx.db, ds), str(tmp_path / "o.pdf")))
    assert "Outstanding fee report" in t and "Asha Verma" in t
    c = export.to_csv(ds, str(tmp_path / "o.csv"))
    assert "Asha Verma" in open(c, encoding="utf-8-sig").read()
    x = export.to_xlsx(ds, str(tmp_path / "o.xlsx"))
    from openpyxl import load_workbook
    ws = load_workbook(x).active
    assert any(isinstance(cell.value, (int, float)) and cell.value > 0 for row in ws.iter_rows(min_row=4) for cell in row)


def test_attendance_sheets(ctx, admin, student, paid, tmp_path):
    cid = ctx.db.val("SELECT course_id FROM enrollments WHERE student_id=?", (student,))
    t, _ = _text(render.to_pdf(docs.attendance_sheet(ctx.db, admin, cid, None, date.today().isoformat()), str(tmp_path / "d.pdf")))
    assert "DAILY ATTENDANCE SHEET" in t and "Asha Verma" in t
    p = render.to_pdf(docs.attendance_sheet(ctx.db, admin, cid, None, date.today().isoformat(), blank_month=True, month="October 2026"), str(tmp_path / "m.pdf"))
    d = fitz.open(p)
    assert d[0].rect.width > d[0].rect.height and "ATTENDANCE REGISTER" in d[0].get_text()


def test_reminder_text_is_plain_text(ctx):
    t = docs.reminder_text(ctx.db, "Asha", [{"label": "Inst 1", "due": 500000, "due_date": "2026-05-01"}])
    assert "Rs. 5,000.00" in t and "01-05-2026" in t
