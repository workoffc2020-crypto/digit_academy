"""CSV (opens in Excel, Hindi safe) and XLSX export of any report Dataset. Money is written as real numbers in rupees."""
import csv
import os
from datetime import datetime

from app.utils import money


def _val(kind, v):
    if v is None:
        return ""
    if kind == "money":
        return money.plain(int(v))
    if kind == "pct":
        return f"{float(v):.2f}"
    if kind == "date":
        try:
            return datetime.fromisoformat(str(v)).strftime("%d-%m-%Y")
        except ValueError:
            return str(v)
    return v


def to_csv(ds, path: str) -> str:
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([ds.title])
        if ds.subtitle:
            w.writerow([ds.subtitle])
        w.writerow([n for n, _ in ds.columns])
        for r in ds.rows:
            w.writerow([_val(k, v) for (n, k), v in zip(ds.columns, r)])
        if ds.totals:
            w.writerow([_val(k, v) for (n, k), v in zip(ds.columns, ds.totals)])
    return path


def to_xlsx(ds, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = (ds.title[:28] or "Report").replace("/", "-")
    ws.append([ds.title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([ds.subtitle])
    ws.append([n for n, _ in ds.columns])
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B2545")
    for r in ds.rows + ([ds.totals] if ds.totals else []):
        row = []
        for (n, k), v in zip(ds.columns, r):
            if k == "money" and v not in (None, ""):
                row.append(int(v) / 100)
            elif k == "int" and v not in (None, ""):
                row.append(int(v))
            elif k == "pct" and v not in (None, ""):
                row.append(float(v))
            else:
                row.append(_val(k, v))
        ws.append(row)
    for j, (n, k) in enumerate(ds.columns, 1):
        for cell in ws.iter_rows(min_row=4, min_col=j, max_col=j):
            if k == "money":
                cell[0].number_format = "#,##0.00"
            if cell[0].value is not None and k in ("text",) and isinstance(cell[0].value, str) and cell[0].value[:1] == "=":
                cell[0].value = "'" + cell[0].value                      # never let a name start a formula
        width = max([len(str(n))] + [len(str(_val(k, r[j - 1]))) for r in ds.rows[:200] if len(r) >= j]) + 2
        ws.column_dimensions[ws.cell(row=3, column=j).column_letter].width = min(max(width, 8), 45)
    ws.freeze_panes = "A4"
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return path
